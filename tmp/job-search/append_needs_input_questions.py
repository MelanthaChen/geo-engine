from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path("/Users/melc/Documents/Intern/Yuxuan_Job_Application_Packet")
PRIMARY = ROOT / "Yuxuan_Chen_Application_Questions.docx"
SECONDARY = ROOT / "05_Portfolio_and_Writing" / "Yuxuan_Chen_Application_Questions.docx"


def add_response_line(document: Document, lines: int = 2) -> None:
    for _ in range(lines):
        p = document.add_paragraph()
        p.add_run("Response: ").bold = True
        p.add_run("_" * 78)


doc = Document(PRIMARY)

doc.add_heading("September 9 2026  Needs input follow up", level=2)
doc.add_paragraph(
    "Only genuinely missing applicant answers are listed here. Items already answered elsewhere in "
    "this document were not repeated. Transcript requests are excluded because required transcripts "
    "mean the role should be skipped, and optional transcripts should not be submitted."
)

doc.add_heading("Canonical  Software Engineer Python Cloud graduate level", level=3)
doc.add_paragraph(
    "Canonical requires these answers to be written entirely in your own words and may disqualify "
    "AI-generated responses. Please answer each prompt with concrete examples you can personally explain."
)
doc.add_paragraph(
    "27. Describe your hands-on experience with Python and Linux, including one project or situation "
    "that best demonstrates your ability."
)
add_response_line(doc, 3)
doc.add_paragraph(
    "28. Describe your experience with cloud provisioning or cloud infrastructure. Name the tools or "
    "platforms you actually used and what you personally implemented."
)
add_response_line(doc, 3)
doc.add_paragraph(
    "29. Describe any real experience you have with Go or Rust. If you have none, write that honestly "
    "and describe the closest relevant systems or backend experience."
)
add_response_line(doc, 3)

doc.add_heading("DefenseStorm  Junior Product Engineer Software Engineer I", level=3)
doc.add_paragraph(
    "30. What percentage of the code you currently write is written by an AI coding agent? Give one "
    "honest percentage from 0 to 100 and, optionally, one sentence explaining how you use the agent."
)
add_response_line(doc, 2)

doc.add_heading("Prometheum  Software Engineer 1 Full Stack", level=3)
doc.add_paragraph(
    "31. Provide two US-based former managers or direct supervisors who may serve as professional "
    "references. For each person include full name, title, company, relationship to you, email, phone, "
    "and whether you have already told them they may be contacted."
)
add_response_line(doc, 5)

doc.add_heading("Existing unanswered items still blocking applications", level=3)
doc.add_paragraph(
    "The following were already asked above and were not duplicated: IMC application history "
    "(Question 7), Amazon or Twitch application history (Question 8), and a recently read technical "
    "paper, blog post, or documentation item (Question 12)."
)

doc.add_heading("Website actions requiring you", level=3)
actions = [
    "Sign in or create the applicant account: Lenovo, Barclays, TikTok or ByteDance, Microsoft, Salesforce, Insperity, Amazon or Audible, Bentley Systems, and T-Mobile.",
    "Complete the DefenseStorm reCAPTCHA after answering Question 30.",
    "If a site asks for a one-time code or account verification, complete it in the browser; never place a password or verification code in this document.",
]
for action in actions:
    doc.add_paragraph(action, style="List Bullet")
doc.add_paragraph(
    "The remaining Needs input entries are browser or uploader failures rather than missing personal "
    "information: LinkedIn phone-field persistence, resume uploaders that do not open, and Jobright "
    "LinkedIn verification. No written answer from you is needed for those."
)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)

doc.save(PRIMARY)
SECONDARY.parent.mkdir(parents=True, exist_ok=True)
doc.save(SECONDARY)

