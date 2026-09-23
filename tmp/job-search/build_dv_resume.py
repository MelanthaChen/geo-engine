from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/melc/Documents/Intern/LLM Search Optimizer/geo-engine/outputs/01a07a13-0189-70c1-a1a4-d8e4c1f7d8e0/Yuxuan_Chen_Resume_DV_Trading_Graduate_Software_Engineer_20260908.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.52)
sec.bottom_margin = Inches(0.52)
sec.left_margin = Inches(0.58)
sec.right_margin = Inches(0.58)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.4)
styles["Normal"].paragraph_format.space_after = Pt(0)
styles["Normal"].paragraph_format.line_spacing = 1.08

def hyperlink(p, text, url):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "000000")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    rp.extend([c, u]); r.append(rp)
    t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r); p._p.append(h)

def section(title):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(10.2); r.font.color.rgb = RGBColor(0,0,0)
    pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "5"); bt.set(qn("w:color"), "555555")
    pb.append(bt); p._p.get_or_add_pPr().append(pb)

def row(left, right):
    table = doc.add_table(rows=1, cols=2); table.autofit = False
    table.columns[0].width = Inches(5.85); table.columns[1].width = Inches(1.35)
    for cell in table.rows[0].cells:
        mar = OxmlElement("w:tcMar")
        for side in ("top","left","bottom","right"):
            e=OxmlElement(f"w:{side}"); e.set(qn("w:w"),"0"); e.set(qn("w:type"),"dxa"); mar.append(e)
        cell._tc.get_or_add_tcPr().append(mar)
    p1,p2=table.cell(0,0).paragraphs[0],table.cell(0,1).paragraphs[0]
    p2.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    a,b=p1.add_run(left),p2.add_run(right); a.bold=True; b.bold=True

def bullet(text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.15); p.paragraph_format.first_line_indent=Inches(-0.12)
    p.paragraph_format.space_after=Pt(1.5); p.add_run("• "); p.add_run(text)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
r=p.add_run("Yuxuan Chen"); r.bold=True; r.font.size=Pt(16); r.font.name="Arial"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
p.add_run("Malden, MA | yuxuan.chen.1031@gmail.com | +1 (223) 269-8058 | ")
hyperlink(p,"LinkedIn","https://www.linkedin.com/in/yuxuan-chen-739160245/"); p.add_run(" | "); hyperlink(p,"GitHub","https://github.com/MelanthaChen")

section("Summary")
p=doc.add_paragraph("Software engineering graduate student with hands-on experience building Python systems in Linux-based cloud environments, profiling concurrent workloads, and delivering data-intensive APIs and automation. Combines computer science and quantitative economics training with end-to-end ownership of testing, deployment, telemetry, and performance analysis.")

section("Education")
row("Northeastern University, Boston, MA", "Expected Dec 2026")
p=doc.add_paragraph("M.S. in Software Engineering Systems"); p.paragraph_format.space_after=Pt(1)
row("Dickinson College, Carlisle, PA", "May 2024")
doc.add_paragraph("B.S. in Computer Science and B.S. in Quantitative Economics")

section("Technical Skills")
p=doc.add_paragraph(); p.add_run("Languages: ").bold=True; p.add_run("Python, Java, TypeScript, SQL")
p=doc.add_paragraph(); p.add_run("Systems and Data: ").bold=True; p.add_run("Linux, FastAPI, REST APIs, PostgreSQL, SQLAlchemy, Concurrent Workloads, Data Pipelines")
p=doc.add_paragraph(); p.add_run("Engineering Tools: ").bold=True; p.add_run("Git, Docker, CI/CD, AWS, GCP, Playwright, React, OpenAI APIs, vLLM")

section("Experience")
row("Research Assistant, Northeastern University - Boston, MA", "May 2026 - Present")
bullet("Architected a full-stack experimentation platform with a React interface and scalable Python and FastAPI services for benchmark evaluation and recommendation analysis.")
bullet("Designed and implemented 47 REST APIs with PostgreSQL and OpenAI integrations, owning workflows from requirements and data modeling through testing and delivery.")
bullet("Built reliable browser automation and persistent-session workflows with Playwright, translating operational requirements into maintainable software components.")
bullet("Implemented a benchmark-driven engine reproducing the Princeton GEO methodology with configurable strategies, experiment aggregation, and reproducible evaluation.")

row("Software Engineer Intern, Software Velocity Corporation - Boston, MA", "Sep 2025 - Jan 2026")
bullet("Built a Python evaluation platform for 10+ language models across Linux-based cloud GPU environments and external APIs, tracking latency, throughput, cost, and configuration metadata.")
bullet("Orchestrated concurrent inference workloads, profiled system bottlenecks, and improved end-to-end pipeline efficiency by 25-40%.")
bullet("Automated cloud deployment, functional evaluation, and execution workflows, reducing setup time and improving reliability of large-scale testing.")
bullet("Developed telemetry and experiment tracking for model, GPU, prompt-set, and runtime behavior to support analytical comparison and production troubleshooting.")

row("Research Assistant, Dickinson College - Carlisle, PA", "Sep 2023 - Jul 2024")
bullet("Built Python NLP and data-processing pipelines for 300K+ documents and 100M+ data points, producing structured corpora for downstream quantitative analysis.")
bullet("Developed reusable querying and visualization components that accelerated iterative analysis of high-dimensional text data.")
bullet("Designed repeatable corpus-generation workflows with clear data transformations, validation, and reusable outputs.")
bullet("Performed linguistic analysis on high-dimensional text data to identify patterns for model-driven insights.")

section("Project")
row("Account Lifecycle Management Platform", "May 2026 - Present")
bullet("Developed a web-based platform with provider orchestration, campaign scheduling, persistent sessions, health monitoring, and 49 REST APIs.")
bullet("Designed modular workflow abstractions and provider-specific components for dependable multi-platform execution and safe extension of an established codebase.")
bullet("Implemented scheduling and recommendation services around persistent sessions, with 49 REST APIs supporting reliable operational workflows.")

doc.save(OUT)
print(OUT)
