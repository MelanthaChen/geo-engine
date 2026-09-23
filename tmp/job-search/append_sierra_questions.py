from docx import Document
from docx.shared import Pt


path = "/Users/melc/Documents/Intern/Yuxuan_Job_Application_Packet/Yuxuan_Chen_Application_Questions.docx"
doc = Document(path)

existing = "\n".join(p.text for p in doc.paragraphs)
if "Sierra  Software Engineer Agent New Grad 2027" not in existing:
    heading = doc.add_paragraph()
    heading.style = "Heading 2"
    heading.add_run("Sierra  Software Engineer Agent New Grad 2027")

    questions = [
        (
            "32. How should your legal name be pronounced? Sierra requires a phonetic spelling. "
            "Write the pronunciation you want recruiters to use for “Yuxuan Chen.”"
        ),
        (
            "33. Sierra offers only Winter 2027, Spring 2027, or Fall 2027 for the required anticipated "
            "graduation season, but your actual graduation date is December 2026. Which option does Sierra "
            "instruct December 2026 graduates to choose? If you contact recruiting, paste their answer here."
        ),
    ]
    for question in questions:
        p = doc.add_paragraph(question)
        p.paragraph_format.space_after = Pt(4)
        for _ in range(3):
            response = doc.add_paragraph("Response: " + "_" * 78)
            response.paragraph_format.space_after = Pt(2)

    doc.save(path)
