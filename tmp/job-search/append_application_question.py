from docx import Document
from docx.shared import Pt

PATH = "/Users/melc/Documents/Intern/Yuxuan_Job_Application_Packet/Yuxuan_Chen_Application_Questions.docx"

doc = Document(PATH)
needle = "28. Prometheum requires two U.S.-based professional references"
if not any(needle in p.text for p in doc.paragraphs):
    doc.add_paragraph("September 8 2026  Prometheum")
    doc.add_paragraph(needle)
    doc.add_paragraph(
        "The Software Engineer 1 Full-Stack application requires two former managers or direct supervisors "
        "who are based in the United States and may be contacted during the hiring process."
    )
    doc.add_paragraph("Reference 1 name: __________________________________________________________")
    doc.add_paragraph("Reference 1 title and organization: _____________________________________________")
    doc.add_paragraph("Reference 1 email: ___________________________________________________________")
    doc.add_paragraph("Reference 1 phone: __________________________________________________________")
    doc.add_paragraph("Reference 1 relationship and dates supervised: __________________________________")
    doc.add_paragraph("Reference 2 name: __________________________________________________________")
    doc.add_paragraph("Reference 2 title and organization: _____________________________________________")
    doc.add_paragraph("Reference 2 email: ___________________________________________________________")
    doc.add_paragraph("Reference 2 phone: __________________________________________________________")
    doc.add_paragraph("Reference 2 relationship and dates supervised: __________________________________")
    doc.add_paragraph(
        "Application status: waiting for these two required references. The form also asks about work "
        "authorization, future sponsorship, desired salary, LinkedIn, portfolio, residence, and background "
        "check; those answers are already available."
    )
    doc.save(PATH)
else:
    start = next(i for i, p in enumerate(doc.paragraphs) if "September 8 2026  Prometheum" in p.text)
    for p in doc.paragraphs[start:]:
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.size = Pt(10)
    doc.save(PATH)

doc = Document(PATH)
uniswap_needle = "29. What blockchain project are you most excited about, and how do you interact with it"
start = next((i for i, p in enumerate(doc.paragraphs) if "September 8 2026  Uniswap Labs" in p.text), None)
if start is not None:
    for paragraph in list(doc.paragraphs[start:]):
        paragraph._element.getparent().remove(paragraph._element)
    doc.save(PATH)
