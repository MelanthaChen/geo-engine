from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/melc/Documents/Intern/LLM Search Optimizer/geo-engine/outputs/01a07a13-0189-70c1-a1a4-d8e4c1f7d8e0/Yuxuan_Chen_Resume_Giga_Software_Engineer_AI_Agents_20260907.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.52)
sec.bottom_margin = Inches(0.52)
sec.left_margin = Inches(0.58)
sec.right_margin = Inches(0.58)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.4)
styles["Normal"].paragraph_format.space_after = Pt(0)
styles["Normal"].paragraph_format.line_spacing = 1.08

def hyperlink(p, text, url):
    part = p.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "000000")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rPr.extend([color, underline]); r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r); p._p.append(h)

def section(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper())
    r.bold = True; r.font.size = Pt(10.2); r.font.color.rgb = RGBColor(0,0,0)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "5"); bottom.set(qn("w:color"), "555555")
    pbdr.append(bottom); pPr.append(pbdr)

def row(left, right, italic=False):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.85); table.columns[1].width = Inches(1.35)
    table.rows[0].cells[0].width = Inches(5.85); table.rows[0].cells[1].width = Inches(1.35)
    for cell in table.rows[0].cells:
        cell.margin_top = 0; cell.margin_bottom = 0
        tcPr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
        for side in ("top","left","bottom","right"):
            e=OxmlElement(f"w:{side}"); e.set(qn("w:w"),"0"); e.set(qn("w:type"),"dxa"); mar.append(e)
        tcPr.append(mar)
    p1=table.cell(0,0).paragraphs[0]; p2=table.cell(0,1).paragraphs[0]
    p1.paragraph_format.space_after=Pt(0); p2.paragraph_format.space_after=Pt(0); p2.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    a=p1.add_run(left); b=p2.add_run(right); a.bold=True; b.bold=True
    if italic: a.italic=True
    return table

def bullet(text):
    p=doc.add_paragraph(style=None)
    p.paragraph_format.left_indent=Inches(0.15); p.paragraph_format.first_line_indent=Inches(-0.12)
    p.paragraph_format.space_after=Pt(1.5); p.add_run("• "); p.add_run(text)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after=Pt(1)
r=p.add_run("Yuxuan Chen"); r.bold=True; r.font.size=Pt(16); r.font.name="Arial"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
p.add_run("Malden, MA | yuxuan.chen.1031@gmail.com | +1 (223) 269-8058 | ")
hyperlink(p,"LinkedIn","https://www.linkedin.com/in/yuxuan-chen-739160245/"); p.add_run(" | "); hyperlink(p,"GitHub","https://github.com/MelanthaChen")

section("Summary")
p=doc.add_paragraph("Software engineering graduate student building production-oriented AI, backend, and full-stack systems with Python, FastAPI, TypeScript, React, PostgreSQL, cloud infrastructure, and LLM workflows. Experienced in end-to-end ownership, performance evaluation, concurrent workloads, and maintainable API design.")

section("Education")
row("Northeastern University, Boston, MA", "Expected Dec 2026")
p=doc.add_paragraph("M.S. in Software Engineering Systems"); p.paragraph_format.space_after=Pt(1)
row("Dickinson College, Carlisle, PA", "May 2024")
doc.add_paragraph("B.S. in Computer Science and B.S. in Quantitative Economics")

section("Technical Skills")
p=doc.add_paragraph(); p.add_run("Languages: ").bold=True; p.add_run("Python, TypeScript, Java, SQL")
p=doc.add_paragraph(); p.add_run("Backend and Data: ").bold=True; p.add_run("FastAPI, REST APIs, PostgreSQL, SQLAlchemy, Data Pipelines, Playwright, React")
p=doc.add_paragraph(); p.add_run("AI Agents and Infrastructure: ").bold=True; p.add_run("LLM Systems, RAG, NLP, OpenAI APIs, vLLM, Telemetry, AWS, GCP, Docker, CI/CD")

section("Experience")
row("Research Assistant, Northeastern University - Boston, MA", "May 2026 - Present")
bullet("Architected a full-stack experimentation platform for Generative Engine Optimization, connecting a React interface to scalable Python and FastAPI services for benchmark evaluation and recommendation analysis.")
bullet("Designed and implemented 47 REST APIs with PostgreSQL and OpenAI integrations to orchestrate AI-agent experiments, evaluation workflows, and persistent results.")
bullet("Built agentic browser workflows with Playwright for reusable session management, multi-platform actions, and reliable content delivery.")
bullet("Implemented a benchmark-driven engine reproducing the Princeton GEO methodology with configurable strategies, experiment aggregation, and reproducible evaluation.")

row("Software Engineer Intern, Software Velocity Corporation - Boston, MA", "Sep 2025 - Jan 2026")
bullet("Built a unified evaluation platform for 10+ language models across cloud GPU environments and external APIs, supporting reliable comparison of latency, throughput, and cost.")
bullet("Developed telemetry and experiment tracking with model, configuration, GPU, and prompt-set metadata, improving visibility into AI-system behavior and operational reliability.")
bullet("Orchestrated concurrent inference workloads, identified system bottlenecks, and improved end-to-end pipeline efficiency by 25-40%.")
bullet("Automated cloud deployment, functional evaluation, and execution workflows, reducing setup time and increasing the reliability of large-scale testing.")

row("Research Assistant, Dickinson College - Carlisle, PA", "Sep 2023 - Jul 2024")
bullet("Built Python NLP and data-processing pipelines for 300K+ documents and 100M+ data points, producing structured corpora for downstream analysis and modeling.")
bullet("Developed reusable querying and visualization components that accelerated iterative analysis of high-dimensional text data.")
bullet("Designed structured corpus-generation workflows that supported repeatable downstream NLP analysis and model-driven insights.")
bullet("Performed linguistic analysis on high-dimensional text data to identify patterns for model-driven insights.")

section("Project")
row("Account Lifecycle Management Platform", "May 2026 - Present")
bullet("Developed a web-based account lifecycle platform with provider-based orchestration, persistent browser sessions, campaign scheduling, and 49 REST APIs.")
bullet("Designed modular workflow abstractions with health monitoring and recommendation services to keep integrations maintainable and extensible.")
bullet("Implemented scheduling and provider-specific workflow components around persistent sessions for dependable multi-platform execution.")

doc.save(OUT)
print(OUT)
