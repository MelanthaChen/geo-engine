from copy import deepcopy
from docx import Document


SOURCE = "/Users/melc/Documents/Intern/Yuxuan_Job_Application_Packet/05_Portfolio_and_Writing/Yuxuan_Chen_Application_Questions.docx"
OUTPUT = "/private/tmp/Yuxuan_Chen_Application_Questions.docx"
HEADING = "September 8 2026  Prophet Security"


doc = Document(SOURCE)
if not any(paragraph.text.strip() == HEADING for paragraph in doc.paragraphs):
    heading = doc.add_paragraph(HEADING)
    heading.style = "Heading 2"
    question = doc.add_paragraph(
        "Why are you interested in the Software Engineer Backend New Graduate role at Prophet Security? "
        "Please answer in a few sentences entirely in your own words. The employer explicitly prohibits "
        "AI assisted application responses."
    )
    question.style = "Normal"
    response = doc.add_paragraph("Response: ____________________________________________________________________")
    response.style = "Normal"
    note = doc.add_paragraph("Do not use an AI written response for this question.")
    note.style = "Normal"

doc.save(OUTPUT)
