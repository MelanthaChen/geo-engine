from docx import Document
from docx.shared import Pt

SOURCE = "/Users/melc/Documents/Intern/Yuxuan_Job_Application_Packet/Yuxuan_Chen_Application_Questions.docx"
OUTPUT = "/Users/melc/Documents/Intern/LLM Search Optimizer/geo-engine/tmp/job-search/Yuxuan_Chen_Application_Questions_updated.docx"

doc = Document(SOURCE)

date_line = doc.add_paragraph()
date_line.style = doc.styles["Heading 2"]
date_line.add_run("September 8 2026  DefenseStorm")

question = doc.add_paragraph()
question.add_run("27. What percentage of your code is written by an agent right now").bold = True

context = doc.add_paragraph(
    "Required for the Junior Product Engineer Software Engineer I application. "
    "Enter your best honest estimate as a percentage."
)
context.paragraph_format.space_after = Pt(3)

response = doc.add_paragraph("Response: ____________________________________________________________________")
response.paragraph_format.space_after = Pt(8)

status = doc.add_paragraph(
    "Application status: all other fields are completed and the matching AI Agents resume is attached. "
    "The remaining website step is the reCAPTCHA human check."
)
status.paragraph_format.space_after = Pt(8)

doc.save(OUTPUT)
