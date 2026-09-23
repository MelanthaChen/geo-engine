from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "/Users/melc/Documents/Intern/LLM Search Optimizer/geo-engine/outputs/01a07a13-0189-70c1-a1a4-d8e4c1f7d8e0/Yuxuan_Chen_Application_Questions.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size in [("Heading 1", 15), ("Heading 2", 12)]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(5)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")

def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True

def add_question(number, prompt, guidance=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number}. {prompt}")
    r.bold = True
    if guidance:
        g = doc.add_paragraph(guidance)
        g.paragraph_format.keep_with_next = True
        g.paragraph_format.left_indent = Inches(0.18)
        g.paragraph_format.space_after = Pt(3)
        g.runs[0].italic = True
        g.runs[0].font.color.rgb = RGBColor(80, 80, 80)
    a = doc.add_paragraph("Response: ")
    a.paragraph_format.keep_together = True
    a.paragraph_format.left_indent = Inches(0.18)
    a.paragraph_format.space_after = Pt(7)
    a.add_run("____________________________________________________________________")

title = doc.add_paragraph(style="Title")
title.add_run("Yuxuan Chen Application Questions")

intro = doc.add_paragraph(
    "This document collects information that only the applicant can confirm or describe. "
    "Responses may be written in English or Chinese. New questions will be appended at the bottom without changing existing answers."
)
intro.paragraph_format.space_after = Pt(9)

meta = doc.add_table(rows=2, cols=3)
meta.autofit = False
widths = [1.35, 2.2, 3.25]
labels = [("Status", "Waiting for applicant", "Last updated September 7 2026"),
          ("How to respond", "Type directly in Word", "Never enter passwords SSN or bank details")]
for i, row in enumerate(meta.rows):
    for j, cell in enumerate(row.cells):
        cell.width = Inches(widths[j])
        cell.text = labels[i][j]
        set_cell_margins(cell)
        set_cell_shading(cell, "EAF2F8" if j == 0 else ("F7F9FB" if i % 2 else "FFFFFF"))
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(9.5)
            run.font.bold = (j == 0)

h = doc.add_paragraph("Work authorization clarification", style="Heading 1")
keep_with_next(h)
add_question(1, "Have you previously used OPT", "You previously stated both OPT pending and prior OPT use. If used, list the degree, start date, and end date. Otherwise answer No.")
add_question(2, "What authorizes your current Northeastern Research Assistant work", "For example F 1 on campus employment, CPT, or another basis. Include the authorization end date.")
add_question(3, "What is the Program End Date on your current I 20")
add_question(4, "What are the expiration dates on your F 1 visa stamp and passport")
add_question(5, "Has the OPT I 765 been submitted", "If submitted, provide the receipt date and current status. The receipt number is optional.")
add_question(6, "May applications state February 18 2027 or earlier if work authorization permits", "Answer Yes or No.")

h = doc.add_paragraph("Previous applications", style="Heading 1")
keep_with_next(h)
add_question(7, "Have you applied to IMC during the past 12 to 18 months", "Answer Yes, No, or Unsure.")
add_question(8, "Have you applied to Amazon Twitch or another Amazon subsidiary", "Answer Yes, No, or Unsure.")
add_question(9, "Before this current search had you applied to Cisco Precisely NiCE or Abax Health", "Answer separately for each company. Do not count applications submitted during this current search.")
add_question(10, "Has any company told you that you may not reapply for a specific period")

h = doc.add_paragraph("Applicant written responses", style="Heading 1")
keep_with_next(h)
add_question(11, "Palantir example of changing your mind", "Describe what you originally believed, what happened, how your view changed, and the result. Maximum 150 words. Chinese is acceptable.")
add_question(12, "A technical paper blog post or documentation you genuinely read in the past month", "Provide the title or link and one thing you learned. If none, answer None.")
add_question(13, "Why do you want to work in software engineering or AI engineering", "Use your own words in two to four sentences.")
add_question(14, "Why are you willing to join an early stage startup", "Use your own words in two to four sentences.")
add_question(15, "A true example that best demonstrates how you solved a difficult problem", "Describe the context, your actions, and the result.")

h = doc.add_paragraph("Travel and role boundaries", style="Heading 1")
keep_with_next(h)
add_question(16, "What state issued your US driver license and when does it expire")
add_question(17, "Do you have a car")
add_question(18, "What is the maximum travel percentage you will accept", "For example 25 percent or 50 percent.")
add_question(19, "Will you accept an out of state role without relocation assistance")
add_question(20, "Will you accept a contract to hire role")
add_question(21, "Accept pay under $120,000 if it meets the local H-1B Level 1 wage")

h = doc.add_paragraph("Documents and work samples", style="Heading 1")
keep_with_next(h)
add_question(22, "Can you provide a Northeastern unofficial transcript")
add_question(23, "Can you provide a Dickinson unofficial transcript")
add_question(24, "Which two or three GitHub repositories should recruiters review")
add_question(25, "Do you have papers posters presentations a portfolio or public writing")
add_question(26, "Do you hold AWS GCP or other certifications")

h = doc.add_paragraph("Website actions requiring the applicant", style="Heading 1")
keep_with_next(h)
p = doc.add_paragraph()
p.add_run("Currently pending: ").bold = True
p.add_run("TikTok login, Insperity Workday login or registration, and Lenovo Workday login or registration. Record the date here after completing any item.")
doc.add_paragraph("Completion record: __________________________________________________________")

h = doc.add_paragraph("New questions appended below", style="Heading 1")
keep_with_next(h)
doc.add_paragraph(
    "All future questions will be appended below this heading and labeled with the discovery date and company. Existing questions and responses will remain in their original order."
)
doc.add_paragraph("September 7 2026  No additional questions")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Yuxuan Chen Job Application Questions")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

def force_cjk_font(run):
    run.font.name = "Hiragino Sans GB"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Hiragino Sans GB")
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        force_cjk_font(run)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    force_cjk_font(run)
for paragraph in section.footer.paragraphs:
    for run in paragraph.runs:
        force_cjk_font(run)

doc.save(OUT)
print(OUT)
